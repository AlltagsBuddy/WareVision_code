"""Lieferanten-Artikeldaten: Datei einlesen (CSV/Excel), Spalten zuordnen, speichern, nachschlagen.

Die Artikeldatei eines Großhändlers (EAN, Bezeichnung, Hersteller, Bestellnummer, EK …)
wird als Katalog gespeichert. Beim Anlegen eines Artikels füllt ein Scan des Barcodes
das Formular aus dem Katalog vor – die Werte werden nur vorgeschlagen und vor dem
Speichern geprüft.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.models.supplier_catalog import SupplierCatalogItem
from app.services.barcode import barcode_variants, clean_barcode, looks_like_gtin

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000

# Zielfelder mit Bezeichnung und typischen Spaltenüberschriften (vereinfacht geschrieben)
FIELDS: dict[str, tuple[str, tuple[str, ...]]] = {
    "barcode": ("Barcode (EAN/GTIN)", ("ean", "gtin", "ean13", "eancode", "barcode", "strichcode", "eangtin",
                                       "gtinean", "eannummer", "upc")),
    "supplier_article_number": ("Artikelnummer beim Lieferanten", (
        "artikelnummer", "artnr", "artikelnr", "bestellnummer", "bestellnr", "lieferantenartikelnummer",
        "itemnumber", "sku", "materialnummer", "artikelno", "art", "nummer", "lieferantennr")),
    "name": ("Bezeichnung", ("bezeichnung", "artikelbezeichnung", "name", "artikelname", "produktname",
                             "kurztext", "bezeichnung1", "titel", "produktbezeichnung", "description")),
    "description": ("Beschreibung (lang)", ("langtext", "beschreibung", "langbeschreibung", "bezeichnung2",
                                            "zusatztext", "details")),
    "manufacturer": ("Hersteller / Marke", ("hersteller", "marke", "brand", "manufacturer", "herstellername",
                                            "fabrikat")),
    "manufacturer_part_number": ("Herstellernummer / OEM", (
        "herstellernummer", "herstellerartikelnummer", "oem", "oemnummer", "mpn", "originalnummer",
        "herstellernr", "teilenummer")),
    "purchase_price": ("Einkaufspreis netto (EK)", ("ek", "einkaufspreis", "ekpreis", "eknetto", "nettopreis",
                                                    "haendlerpreis", "ihrpreis", "hek", "netto")),
    "retail_price_gross": ("Verkaufspreis/UVP inkl. MwSt", ("uvp", "vkbrutto", "verkaufspreisbrutto",
                                                           "uvpbrutto", "ladenpreis", "brutto", "endkundenpreis")),
    "retail_price_net": ("Verkaufspreis netto", ("vknetto", "uvpnetto", "verkaufspreisnetto", "vk",
                                                 "verkaufspreis")),
    "vat_rate": ("MwSt-Satz (%)", ("mwst", "ust", "steuersatz", "mwstsatz", "ustsatz", "mehrwertsteuer")),
    "unit": ("Einheit", ("einheit", "me", "mengeneinheit", "unit", "ve")),
}


# Allgemeine Wörter nur bei exakt gleicher Überschrift verwenden (sonst falsche Treffer,
# z. B. „nummer“ in „Herstellernummer“ oder „netto“ in „VK netto“)
EXACT_ONLY = {"nummer", "netto", "brutto", "art", "vk", "me", "ve", "name", "preis", "ek", "ust", "mwst", "oem",
              "upc", "sku", "titel", "verkaufspreis", "description"}


class CatalogError(ValueError):
    """Fehler beim Einlesen (Meldung ist für den Nutzer gedacht)."""


def _norm(text: str) -> str:
    t = (text or "").strip().lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        t = t.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", t)


def suggest_mapping(headers: list[str]) -> dict[str, int | None]:
    """Spalten anhand der Überschriften vorschlagen (exakter Treffer vor Teiltreffer)."""
    normed = [_norm(h) for h in headers]
    used: set[int] = set()
    mapping: dict[str, int | None] = {}
    for key, (_, synonyms) in FIELDS.items():
        pick = None
        for syn in synonyms:
            for i, h in enumerate(normed):
                if i not in used and h == syn:
                    pick = i
                    break
            if pick is not None:
                break
        if pick is None:
            for syn in synonyms:
                if len(syn) < 4 or syn in EXACT_ONLY:
                    continue
                for i, h in enumerate(normed):
                    if i not in used and syn in h:
                        pick = i
                        break
                if pick is not None:
                    break
        mapping[key] = pick
        if pick is not None:
            used.add(pick)
    return mapping


def parse_decimal(value) -> Decimal | None:
    """Preis lesen: „1.234,56“, „12,5 €“, „12.50“ oder Zahl aus Excel."""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"), ROUND_HALF_UP)
        except InvalidOperation:
            return None
    s = re.sub(r"[^\d,.\-]", "", str(value))
    if not s or s in ("-", ",", "."):
        return None
    if "," in s and "." in s:
        # Das zuletzt stehende Zeichen ist das Dezimaltrennzeichen
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s).quantize(Decimal("0.01"), ROUND_HALF_UP)
    except InvalidOperation:
        return None


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel speichert EAN oft als Zahl: 4006381333931.0 → 4006381333931
        return str(int(value))
    return str(value).strip()


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]] = field(default_factory=list)


def read_table(content: bytes, filename: str) -> Table:
    """CSV (; , Tab |, UTF-8 oder Windows-Zeichensatz) oder Excel (.xlsx) einlesen."""
    if len(content) > MAX_FILE_BYTES:
        raise CatalogError("Die Datei ist zu groß (höchstens 20 MB).")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        return _read_xlsx(content)
    if name.endswith(".xls"):
        raise CatalogError("Altes Excel-Format (.xls): bitte in Excel als .xlsx oder .csv speichern.")
    return _read_csv(content)


def _read_csv(content: bytes) -> Table:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover – cp1252 dekodiert fast alles
        raise CatalogError("Zeichensatz der Datei nicht erkannt.")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = max(";,\t|", key=lambda d: sample.splitlines()[0].count(d) if sample else 0)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        raise CatalogError("Die Datei ist leer.")
    if len(rows) - 1 > MAX_ROWS:
        raise CatalogError(f"Zu viele Zeilen (höchstens {MAX_ROWS:,}).".replace(",", "."))
    headers = [h.strip() for h in rows[0]]
    return Table(headers=headers, rows=[[c.strip() for c in r] for r in rows[1:]])


def _read_xlsx(content: bytes) -> Table:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 – kaputte Datei → klare Meldung
        raise CatalogError("Die Excel-Datei konnte nicht gelesen werden.") from exc
    ws = wb.worksheets[0]
    headers: list[str] | None = None
    rows: list[list[str]] = []
    for raw in ws.iter_rows(values_only=True):
        cells = [_cell_text(v) for v in raw]
        if not any(cells):
            continue
        if headers is None:
            headers = cells
            continue
        rows.append(cells)
        if len(rows) > MAX_ROWS:
            raise CatalogError(f"Zu viele Zeilen (höchstens {MAX_ROWS:,}).".replace(",", "."))
    wb.close()
    if headers is None:
        raise CatalogError("Die Excel-Datei ist leer.")
    return Table(headers=headers, rows=rows)


def _get(row: list[str], idx: int | None) -> str:
    if idx is None or idx < 0 or idx >= len(row):
        return ""
    return (row[idx] or "").strip()


def _vat(value: str) -> Decimal | None:
    d = parse_decimal(value)
    if d is None:
        return None
    if d < 1 and d > 0:  # 0,19 → 19
        d = d * 100
    return d.quantize(Decimal("0.01"))


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


def import_table(
    db: Session,
    table: Table,
    mapping: dict[str, int | None],
    supplier_name: str,
    supplier_id: uuid.UUID | None,
    source_ref: str,
    default_vat: Decimal = Decimal("19"),
) -> ImportResult:
    """Zeilen speichern. Gleiche Bestellnummer (oder ohne Nummer: gleicher Barcode) desselben
    Lieferanten wird aktualisiert, sonst neu angelegt."""
    if mapping.get("name") is None:
        raise CatalogError("Bitte die Spalte für die Bezeichnung zuordnen.")
    if mapping.get("barcode") is None and mapping.get("supplier_article_number") is None:
        raise CatalogError("Bitte mindestens Barcode oder Artikelnummer zuordnen.")

    existing = db.query(SupplierCatalogItem).filter(SupplierCatalogItem.supplier_name == supplier_name).all()
    by_number = {e.supplier_article_number: e for e in existing if e.supplier_article_number}
    by_barcode = {e.barcode: e for e in existing if e.barcode and not e.supplier_article_number}

    res = ImportResult()
    for line_no, row in enumerate(table.rows, start=2):
        name = _get(row, mapping.get("name"))
        number = _get(row, mapping.get("supplier_article_number")) or None
        barcode = clean_barcode(_get(row, mapping.get("barcode")))
        if barcode and barcode.isdigit() and len(barcode) > 14:
            barcode = None  # z. B. Excel-Rundung, keine gültige GTIN
        if not name or not (number or barcode):
            res.skipped += 1
            continue
        vat = _vat(_get(row, mapping.get("vat_rate"))) or default_vat
        retail_net = parse_decimal(_get(row, mapping.get("retail_price_net")))
        retail_gross = parse_decimal(_get(row, mapping.get("retail_price_gross")))
        if retail_net is None and retail_gross is not None:
            retail_net = (retail_gross / (1 + vat / 100)).quantize(Decimal("0.01"), ROUND_HALF_UP)
        values = {
            "barcode": barcode,
            "supplier_article_number": number,
            "name": name[:500],
            "description": _get(row, mapping.get("description")) or None,
            "manufacturer": (_get(row, mapping.get("manufacturer")) or None),
            "manufacturer_part_number": (_get(row, mapping.get("manufacturer_part_number")) or None),
            "purchase_price": parse_decimal(_get(row, mapping.get("purchase_price"))),
            "retail_price_net": retail_net,
            "vat_rate": vat,
            "unit": (_get(row, mapping.get("unit")) or None),
        }
        for k in ("manufacturer", "manufacturer_part_number"):
            if values[k]:
                values[k] = values[k][:150 if k == "manufacturer" else 100]
        if values["supplier_article_number"]:
            values["supplier_article_number"] = values["supplier_article_number"][:100]
        if values["unit"]:
            values["unit"] = values["unit"][:30]
        item = by_number.get(number) if number else by_barcode.get(barcode)
        if item:
            for k, v in values.items():
                setattr(item, k, v)
            item.source_ref = source_ref
            if supplier_id:
                item.supplier_id = supplier_id
            res.updated += 1
        else:
            item = SupplierCatalogItem(supplier_name=supplier_name, supplier_id=supplier_id, source="import",
                                       source_ref=source_ref, **values)
            db.add(item)
            if number:
                by_number[number] = item
            else:
                by_barcode[barcode] = item
            res.created += 1
    return res


def add_invoice_items(
    db: Session, items: list[dict], supplier_name: str, supplier_id: uuid.UUID | None, invoice_number: str,
) -> ImportResult:
    """Positionen einer Lieferanten-E-Rechnung übernehmen (nur mit Barcode oder Bestellnummer)."""
    res = ImportResult()
    existing = db.query(SupplierCatalogItem).filter(SupplierCatalogItem.supplier_name == supplier_name).all()
    by_number = {e.supplier_article_number: e for e in existing if e.supplier_article_number}
    by_barcode = {e.barcode: e for e in existing if e.barcode}
    for it in items:
        name = (it.get("description") or "").strip()
        number = (it.get("seller_article_number") or "").strip() or None
        barcode = clean_barcode(it.get("global_id")) if looks_like_gtin(it.get("global_id")) else None
        if not name or not (number or barcode):
            res.skipped += 1
            continue
        price = parse_decimal(it.get("unit_price"))
        vat = parse_decimal(it.get("vat_rate"))
        item = (by_number.get(number) if number else None) or (by_barcode.get(barcode) if barcode else None)
        if item:
            # Bestehende Angaben aus der Artikeldatei nicht überschreiben – nur ergänzen bzw. Preis aktualisieren
            item.barcode = item.barcode or barcode
            item.supplier_article_number = item.supplier_article_number or number
            if price is not None:
                item.purchase_price = price
            res.updated += 1
            continue
        item = SupplierCatalogItem(
            supplier_name=supplier_name, supplier_id=supplier_id, source="einvoice", source_ref=invoice_number,
            barcode=barcode, supplier_article_number=number[:100] if number else None, name=name[:500],
            manufacturer=(it.get("brand") or None), purchase_price=price, vat_rate=vat,
            unit=(it.get("unit") or None),
        )
        db.add(item)
        if number:
            by_number[number] = item
        if barcode:
            by_barcode[barcode] = item
        res.created += 1
    return res


def lookup_barcode(db: Session, code: str) -> list[SupplierCatalogItem]:
    """Alle Katalogeinträge zu einem Barcode (neueste zuerst)."""
    variants = barcode_variants(code)
    if not variants:
        return []
    return (db.query(SupplierCatalogItem)
            .filter(SupplierCatalogItem.barcode.in_(variants))
            .order_by(SupplierCatalogItem.updated_at.desc())
            .limit(10).all())


def search(db: Session, q: str, limit: int = 20) -> list[SupplierCatalogItem]:
    """Suche nach Bezeichnung, Bestellnummer, Herstellernummer oder Barcode."""
    term = f"%{q.strip()}%"
    return (db.query(SupplierCatalogItem)
            .filter(or_(SupplierCatalogItem.name.ilike(term),
                        SupplierCatalogItem.supplier_article_number.ilike(term),
                        SupplierCatalogItem.manufacturer_part_number.ilike(term),
                        SupplierCatalogItem.barcode.ilike(term)))
            .order_by(SupplierCatalogItem.name)
            .limit(limit).all())
